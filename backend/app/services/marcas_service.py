from sqlalchemy.orm import Session
from sqlalchemy import func, text
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import List, Optional
import io
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PILImage
from app.models.marcas import Marca, Cuenta
from app.models.inventario import Kardex, Producto
from app.models.ingresos import Ingreso, IngresoDetalle
from app.schemas.marcas import MarcaCreate, MarcaUpdate
from app.core.storage import download_image

class MarcasService:
    @staticmethod
    def get_all(db: Session) -> List[dict]:
        marcas = db.query(Marca).order_by(Marca.id.desc()).all()
        if not marcas:
            return []
        cuentas_por_marca = {}
        for cuenta in db.query(Cuenta).filter(Cuenta.marca_id.in_([m.id for m in marcas])).all():
            cuentas_por_marca.setdefault(cuenta.marca_id, cuenta)
        sku_counts = dict(db.execute(text(
            "SELECT marca_id, COUNT(*) FROM productos GROUP BY marca_id"
        )).all())
        sales_by_brand = dict(db.execute(text("""
            SELECT p.marca_id, COALESCE(SUM(vd.subtotal), 0)
            FROM venta_detalles vd
            JOIN productos p ON vd.producto_id = p.id
            JOIN ventas v ON vd.venta_id = v.id
            WHERE v.fecha_venta >= date_trunc('month', CURRENT_DATE)
              AND v.fecha_venta < date_trunc('month', CURRENT_DATE) + interval '1 month'
            GROUP BY p.marca_id
        """)).all())

        results = []
        for m in marcas:
            cuenta = cuentas_por_marca.get(m.id)
            banco = cuenta.banco if cuenta else None
            no_cuenta = cuenta.no_cuenta if cuenta else None
            sku_count = sku_counts.get(m.id, 0)
            sales_sum = sales_by_brand.get(m.id, 0.0)

            results.append({
                "id": m.id,
                "nombre_marca": m.nombre_marca,
                "cedula": m.cedula or "",
                "codigo": m.codigo or "",
                "categoria": m.categoria or "",
                "razon_social": m.razon_social or "",
                "ciudad_pais": m.ciudad_pais or "",
                "redes_sociales": m.redes_sociales or "",
                "nombre_representante": m.nombre_representante or "",
                "cargo": m.cargo or "",
                "contacto_nombre": m.contacto_nombre or "",
                "telefono": m.telefono or "",
                "email": m.email or "",
                "tiempo_estancia": m.tiempo_estancia or 6,
                "fecha_ingreso": m.fecha_ingreso.isoformat() if m.fecha_ingreso else None,
                "catalogo": m.catalogo or "",
                "estado_id": m.estado_id,
                "logo": m.logo or "",
                "banco": banco,
                "no_cuenta": no_cuenta,
                "activeSkus": f"{sku_count} SKUs" if m.estado_id == 2 else f"{sku_count} SKUs (En pausa)",
                "monthSales": float(sales_sum)
            })
        return results

    @staticmethod
    def get_by_id(db: Session, marca_id: int) -> Optional[dict]:
        m = db.query(Marca).filter(Marca.id == marca_id).first()
        if not m:
            return None
        cuenta = db.query(Cuenta).filter(Cuenta.marca_id == m.id).first()
        banco = cuenta.banco if cuenta else None
        no_cuenta = cuenta.no_cuenta if cuenta else None

        sku_count = db.execute(
            text("SELECT COUNT(*) FROM productos WHERE marca_id = :marca_id"),
            {"marca_id": m.id}
        ).scalar() or 0

        sales_sum = db.execute(
            text("""
                SELECT COALESCE(SUM(vd.subtotal), 0)
                FROM venta_detalles vd
                JOIN productos p ON vd.producto_id = p.id
                JOIN ventas v ON vd.venta_id = v.id
                WHERE p.marca_id = :marca_id
                AND v.fecha_venta >= date_trunc('month', CURRENT_DATE)
                AND v.fecha_venta < date_trunc('month', CURRENT_DATE) + interval '1 month'
            """),
            {"marca_id": m.id}
        ).scalar() or 0.0

        return {
            "id": m.id,
            "nombre_marca": m.nombre_marca,
            "cedula": m.cedula or "",
            "codigo": m.codigo or "",
            "categoria": m.categoria or "",
            "razon_social": m.razon_social or "",
            "ciudad_pais": m.ciudad_pais or "",
            "redes_sociales": m.redes_sociales or "",
            "nombre_representante": m.nombre_representante or "",
            "cargo": m.cargo or "",
            "contacto_nombre": m.contacto_nombre or "",
            "telefono": m.telefono or "",
            "email": m.email or "",
            "tiempo_estancia": m.tiempo_estancia or 6,
            "fecha_ingreso": m.fecha_ingreso.isoformat() if m.fecha_ingreso else None,
            "catalogo": m.catalogo or "",
            "estado_id": m.estado_id,
            "logo": m.logo or "",
            "banco": banco,
            "no_cuenta": no_cuenta,
            "activeSkus": f"{sku_count} SKUs",
            "monthSales": float(sales_sum)
        }

    @staticmethod
    def create(db: Session, marca_in: MarcaCreate) -> dict:
        marca = Marca(
            nombre_marca=marca_in.nombre_marca,
            cedula=marca_in.cedula,
            codigo=marca_in.codigo or (marca_in.nombre_marca[:3].upper() if marca_in.nombre_marca else "MC"),
            categoria=marca_in.categoria,
            razon_social=marca_in.razon_social,
            ciudad_pais=marca_in.ciudad_pais,
            redes_sociales=marca_in.redes_sociales,
            nombre_representante=marca_in.nombre_representante,
            cargo=marca_in.cargo,
            contacto_nombre=marca_in.contacto_nombre,
            telefono=marca_in.telefono,
            email=marca_in.email,
            tiempo_estancia=marca_in.tiempo_estancia,
            fecha_ingreso=marca_in.fecha_ingreso or datetime.utcnow(),
            catalogo=marca_in.catalogo,
            estado_id=marca_in.estado_id,
            logo=marca_in.logo,
            fecha_registro=datetime.utcnow()
        )
        db.add(marca)
        db.commit()
        db.refresh(marca)

        # Crear cuenta bancaria si fue provista
        if marca_in.banco or marca_in.no_cuenta:
            cuenta = Cuenta(
                marca_id=marca.id,
                banco=marca_in.banco,
                no_cuenta=marca_in.no_cuenta
            )
            db.add(cuenta)
            db.commit()
            db.refresh(cuenta)
            marca.cuentas = cuenta.id
            db.commit()

        return MarcasService.get_by_id(db, marca.id)

    @staticmethod
    def update(db: Session, marca_id: int, marca_in: MarcaUpdate) -> Optional[dict]:
        marca = db.query(Marca).filter(Marca.id == marca_id).first()
        if not marca:
            return None

        update_data = marca_in.dict(exclude_unset=True)
        banco = update_data.pop("banco", None)
        no_cuenta = update_data.pop("no_cuenta", None)

        for key, value in update_data.items():
            setattr(marca, key, value)

        db.commit()

        # Actualizar cuenta bancaria
        if banco is not None or no_cuenta is not None:
            cuenta = db.query(Cuenta).filter(Cuenta.marca_id == marca.id).first()
            if cuenta:
                if banco is not None:
                    cuenta.banco = banco
                if no_cuenta is not None:
                    cuenta.no_cuenta = no_cuenta
            else:
                cuenta = Cuenta(
                    marca_id=marca.id,
                    banco=banco or "",
                    no_cuenta=no_cuenta or ""
                )
                db.add(cuenta)
                db.flush()
                marca.cuentas = cuenta.id
            db.commit()

        return MarcasService.get_by_id(db, marca.id)

    @staticmethod
    def update_logo(db: Session, marca_id: int, logo_url: str) -> Optional[dict]:
        marca = db.query(Marca).filter(Marca.id == marca_id).first()
        if not marca:
            return None
        marca.logo = logo_url
        db.commit()
        return MarcasService.get_by_id(db, marca.id)

    @staticmethod
    def generate_delivery_sheet(db: Session, marca_id: int):
        marca = db.query(Marca).filter(Marca.id == marca_id).first()
        if not marca:
            return None, None

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Plantilla_Inventario"
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A11"

        # Paleta de la tienda: rosa, blush, dorado y blanco.
        rose = "E9B7C3"
        pink = "F7D1D6"
        blush = "F3E2D9"
        gold = "C9A46D"
        white = "FFFFFF"
        ink = "493A35"
        muted = "786963"
        line = "E8D8D1"
        zebra = "FFFAF8"
        input_fill = "FFF3F4"
        gold_light = "F8F0E4"
        border = Border(*(Side(style="thin", color=line) for _ in range(4)))
        widths = {"A": 7, "B": 20, "C": 32, "D": 38, "E": 20,
                  "F": 18, "G": 20, "H": 19, "I": 18, "J": 38}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.row_dimensions[1].height = 10
        ws.row_dimensions[2].height = 36
        ws.row_dimensions[3].height = 24
        ws.row_dimensions[4].height = 9
        ws.merge_cells("A2:A3")
        logo_inserted = False
        if getattr(marca, "logo", None):
            logo_content = download_image(marca.logo)
            if logo_content is None:
                clean_logo = marca.logo.lstrip("/").replace("/", os.sep)
                if os.path.exists(clean_logo):
                    with open(clean_logo, "rb") as logo_file:
                        logo_content = logo_file.read()
            if logo_content:
                try:
                    with PILImage.open(io.BytesIO(logo_content)) as image:
                        w, h = image.size
                        ratio = min(82 / w, 50 / h, 1.0)
                    logo = OpenpyxlImage(io.BytesIO(logo_content))
                    logo.width, logo.height = int(w * ratio), int(h * ratio)
                    ws.add_image(logo, "A2")
                    logo_inserted = True
                except Exception:
                    pass
        if not logo_inserted:
            cell = ws["A2"]
            cell.value = (marca.codigo[:3] if marca.codigo else marca.nombre_marca[:3]).upper()
            cell.font = Font(name="Segoe UI", size=14, bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=gold)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("B2:G2")
        ws["B2"] = f"CATÁLOGO DE PRODUCTOS · {marca.nombre_marca.upper()}"
        ws["B2"].font = Font(name="Segoe UI", size=16, bold=True, color=ink)
        ws["B2"].alignment = Alignment(vertical="center")
        ws.merge_cells("B3:G3")
        ws["B3"] = "Plantilla de alta de productos y recepción de stock inicial"
        ws["B3"].font = Font(name="Segoe UI", size=10, color=muted)
        ws["B3"].alignment = Alignment(vertical="center")
        ws.merge_cells("H2:J2")
        ws["H2"] = "IMPORTACIÓN PERSONALIZADA"
        ws["H2"].font = Font(name="Segoe UI", size=9, bold=True, color=ink)
        ws["H2"].fill = PatternFill("solid", fgColor=blush)
        ws["H2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells("H3:J3")
        ws["H3"] = f"Código {marca.codigo or '—'}  —  {datetime.now().strftime('%d/%m/%Y')}"
        ws["H3"].font = Font(name="Segoe UI", size=9, color=muted)
        ws["H3"].fill = PatternFill("solid", fgColor=gold_light)
        ws["H3"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A5:J5")
        ws["A5"] = "DATOS DE LA MARCA"
        ws["A5"].font = Font(name="Segoe UI", size=10, bold=True, color=ink)
        ws["A5"].fill = PatternFill("solid", fgColor=pink)
        ws["A5"].alignment = Alignment(vertical="center")
        for col in range(1, 11):
            ws.cell(5, col).border = border
        details = [
            ("A6", "Marca"), ("B6", marca.nombre_marca), ("D6", "Código"), ("E6", marca.codigo or "—"),
            ("G6", "Categoría"), ("H6", marca.categoria or "General"),
            ("A7", "Representante"), ("B7", marca.nombre_representante or "—"),
            ("D7", "Correo"), ("E7", marca.email or "—"), ("G7", "Teléfono"), ("H7", marca.telefono or "—"),
        ]
        for address, value in details:
            c = ws[address]
            c.value = value
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.font = Font(name="Segoe UI", size=9, bold=address[0] in "ADG", color=ink if address[0] in "ADG" else muted)
            if address[0] in "ADG":
                c.fill = PatternFill("solid", fgColor=gold_light)
        for row in (6, 7):
            ws.row_dimensions[row].height = 25
        ws.merge_cells("B6:C6")
        ws.merge_cells("E6:F6")
        ws.merge_cells("H6:J6")
        ws.merge_cells("B7:C7")
        ws.merge_cells("E7:F7")
        ws.merge_cells("H7:J7")

        ws.merge_cells("A9:J9")
        ws["A9"] = "Completa SKU y cantidad recibida. Para productos nuevos también se necesita nombre y precio de venta."
        ws["A9"].font = Font(name="Segoe UI", size=9, bold=True, color=ink)
        ws["A9"].fill = PatternFill("solid", fgColor=blush)
        ws["A9"].alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[9].height = 25

        headers = ["#", "SKU / CÓDIGO *", "NOMBRE * SI ES NUEVO", "DESCRIPCIÓN", "TALLA / VARIANTE",
                   "PRECIO VENTA C$ * SI ES NUEVO", "CANTIDAD RECIBIDA *", "COSTO UNIT. C$", "STOCK CRÍTICO", "URL DE IMAGEN"]
        for col, title in enumerate(headers, start=1):
            c = ws.cell(10, col, title)
            c.font = Font(name="Segoe UI", size=9, bold=True, color=white)
            c.fill = PatternFill("solid", fgColor=gold if col in (6, 7) else rose)
            c.border = border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.comment = Comment({
                2: "Obligatorio y único por fila. Si ya existe en esta marca, la cantidad recibida se suma al stock. Si pertenece a otra marca, no se importa.",
                3: "Obligatorio para productos nuevos. Si el SKU ya existe en esta marca, se conservan los datos del producto.",
                4: "Detalle útil: material, color, modelo o contenido del paquete.",
                5: "Talla, tamaño o variante; se guarda en el campo de variante del producto.",
                6: "Obligatorio para productos nuevos. Precio de venta en cordobas nicaraguenses; admite decimales, sin simbolo.",
                7: "Obligatorio. Cantidad física recibida, entero mayor que cero. Se suma al stock actual del producto.",
                8: "Opcional. Costo unitario en cordobas nicaraguenses, sin simbolo.",
                9: "Opcional. Umbral para alertas de stock bajo; si se deja vacío, se usa 5.",
                10: "Opcional. URL pública de la imagen del producto.",
            }.get(col, "Número consecutivo automático; no editar."), "Sistema")
        ws.row_dimensions[10].height = 38

        start_row, end_row = 11, 40
        for idx, r in enumerate(range(start_row, end_row + 1), start=1):
            ws.row_dimensions[r].height = 24
            ws.cell(r, 1, idx)
            for col in range(1, 11):
                c = ws.cell(r, col)
                c.border = border
                c.alignment = Alignment(vertical="center", wrap_text=col in (3, 4, 10))
                c.font = Font(name="Segoe UI", size=9, color=muted if col == 1 else ink)
                c.fill = PatternFill("solid", fgColor=(input_fill if col in range(2, 11) else (white if idx % 2 else zebra)))
            for col in (6, 8):
                ws.cell(r, col).number_format = '"C$"#,##0.00'
            ws.cell(r, 7).number_format = "0"
            ws.cell(r, 9, 5)
            ws.cell(r, 9).number_format = "0"

        ws.conditional_formatting.add(
            f"B{start_row}:B{end_row}",
            FormulaRule(
                formula=[f'AND(B{start_row}<>"",COUNTIF($B${start_row}:$B${end_row},B{start_row})>1)'],
                fill=PatternFill("solid", fgColor=rose),
                font=Font(color="9F1239", bold=True),
            ),
        )

        ws.auto_filter.ref = f"A10:J{end_row}"
        dv_text = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="100", allow_blank=True)
        dv_name = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="255", allow_blank=True)
        # El precio es obligatorio para productos nuevos; los existentes solo
        # requieren SKU y cantidad para recibir más unidades.
        dv_price = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
        dv_cost = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        dv_qty = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=False)
        dv_critical = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        for dv in (dv_text, dv_name, dv_price, dv_cost, dv_qty, dv_critical):
            dv.errorTitle = "Dato no válido"
            dv.error = "Revisa el valor y las instrucciones de la plantilla."
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
        dv_text.add(f"B{start_row}:B{end_row}")
        dv_name.add(f"C{start_row}:C{end_row}")
        dv_price.add(f"F{start_row}:F{end_row}")
        dv_qty.add(f"G{start_row}:G{end_row}")
        dv_cost.add(f"H{start_row}:H{end_row}")
        dv_critical.add(f"I{start_row}:I{end_row}")
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
        ws.print_title_rows = "1:10"

        guide = wb.create_sheet("Instrucciones_Carga")
        guide.sheet_view.showGridLines = False
        guide.column_dimensions["A"].width = 5
        guide.column_dimensions["B"].width = 29
        guide.column_dimensions["C"].width = 100
        guide.merge_cells("B2:C2")
        guide["B2"] = f"GUÍA DE CARGA — {marca.nombre_marca.upper()}"
        guide["B2"].font = Font(name="Segoe UI", size=15, bold=True, color=ink)
        guide["B2"].fill = PatternFill("solid", fgColor=pink)
        guide["B2"].alignment = Alignment(vertical="center")
        guide.row_dimensions[2].height = 34
        guide_data = [
            ("Objetivo", "Crea productos nuevos y registra las unidades recibidas. Si el SKU ya existe en esta marca, suma las unidades al stock y conserva sus datos."),
            ("SKU / Código *", "Escribe el SKU del producto. Si ya existe en esta marca, la cantidad se suma al stock; si pertenece a otra marca, la carga se rechaza."),
            ("Nombre *", "Obligatorio para productos nuevos. Para un SKU existente de esta marca puedes dejarlo vacío; los datos actuales del producto se conservan."),
            ("Descripción", "Agrega material, color, contenido del paquete, modelo u otra información que ayude a distinguirlo."),
            ("Talla / Variante", "Talla o variante del producto (por ejemplo: S, 10 cm, color rosa)."),
            ("Precio venta en C$ *", "Obligatorio para productos nuevos. Para un producto existente puedes dejarlo vacío; el precio actual se conserva."),
            ("Cantidad recibida *", "Cantidad física entregada, como entero mayor que cero. En productos existentes se suma al stock actual."),
            ("Costo unitario en C$", "Opcional. Costo por unidad en cordobas nicaraguenses, solo numero decimal. No escribas moneda ni separador de miles."),
            ("Stock crítico", "Opcional. Umbral de alerta de stock bajo; el valor predeterminado es 5."),
            ("URL de imagen", "Opcional. Debe ser un enlace público a la imagen del artículo."),
            ("Antes de importar", "Completa una fila por SKU y cantidad recibida. No cambies el nombre de las hojas, encabezados ni estructura. Elimina filas de ejemplo vacías si agregaste datos fuera de las filas preparadas."),
            ("Resultado", "El sistema valida todo el archivo antes de guardar. Los SKU existentes de esta marca reciben más stock; los nuevos se crean. Si un SKU pertenece a otra marca o hay datos inválidos, no se aplica la carga."),
            ("Campos automáticos", "La marca, el identificador, el código QR y el estado activo los asigna el sistema. No hace falta escribirlos en esta plantilla."),
        ]
        for r, (label, desc) in enumerate(guide_data, start=4):
            guide.row_dimensions[r].height = 42
            a, b = guide.cell(r, 2, label), guide.cell(r, 3, desc)
            a.font = Font(name="Segoe UI", size=9, bold=True, color=ink)
            a.fill = PatternFill("solid", fgColor=blush if r % 2 == 0 else gold_light)
            b.font = Font(name="Segoe UI", size=9, color=ink)
            b.fill = PatternFill("solid", fgColor=white if r % 2 == 0 else zebra)
            for c in (a, b):
                c.border = border
                c.alignment = Alignment(vertical="center", wrap_text=True)

        meta = wb.create_sheet("_plantilla")
        for row, (key, value) in enumerate([
            ("marca_id", marca.id), ("marca_codigo", marca.codigo or ""), ("version", 2),
            ("data_start", start_row), ("data_end", end_row),
        ], start=1):
            meta.cell(row, 1, key)
            meta.cell(row, 2, value)
        meta.sheet_state = "hidden"
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"Plantilla_Productos_{marca.codigo or 'MARCA'}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return output, filename

    @staticmethod
    def import_delivery_template(
        db: Session,
        marca_id: int,
        file_bytes: bytes,
        usuario_id: int,
        filename: Optional[str] = None,
    ) -> dict:
        marca = db.query(Marca).filter(Marca.id == marca_id).first()
        if not marca:
            raise ValueError("Marca no encontrada")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            raise ValueError("El archivo no es una plantilla Excel válida (.xlsx)") from exc
        if "Plantilla_Inventario" not in wb.sheetnames or "_plantilla" not in wb.sheetnames:
            raise ValueError("La plantilla no corresponde al formato de importación de productos")
        meta = wb["_plantilla"]
        if meta["A1"].value != "marca_id" or str(meta["B1"].value) != str(marca.id) or meta["B3"].value != 2:
            raise ValueError("La plantilla pertenece a otra marca o tiene una versión no compatible. Descarga una plantilla nueva.")

        ws = wb["Plantilla_Inventario"]
        data_start, data_end = meta["B4"].value, meta["B5"].value
        if not isinstance(data_start, int) or not isinstance(data_end, int) or data_start < 11 or data_end < data_start:
            raise ValueError("La plantilla no contiene un rango de productos válido")
        rows, seen_skus = [], set()

        def parse_amount(value, label, row_num, required=False, allow_zero=True):
            if value is None or str(value).strip() == "":
                if required:
                    raise ValueError(f"La fila {row_num} necesita {label}")
                return None
            try:
                raw = str(value).strip().replace("C$", "").replace("$", "")
                if isinstance(value, str) and "," in raw:
                    if "." in raw and raw.rfind(",") > raw.rfind("."):
                        raw = raw.replace(".", "").replace(",", ".")
                    elif "." in raw:
                        raw = raw.replace(",", "")
                    else:
                        raw = raw.replace(",", ".")
                amount = Decimal(raw)
            except (InvalidOperation, ValueError) as exc:
                raise ValueError(f"{label} no válido en la fila {row_num}") from exc
            if not amount.is_finite() or amount < 0 or (not allow_zero and amount == 0):
                raise ValueError(f"{label} debe ser {'mayor que cero' if not allow_zero else 'cero o mayor'} (fila {row_num})")
            return amount

        for row_num, values in enumerate(ws.iter_rows(min_row=data_start, max_row=data_end, min_col=1, max_col=10, values_only=True), start=data_start):
            # Ignorar solo filas completamente vacías en las columnas editables.
            editable = values[1:10]
            if all(value is None or str(value).strip() == "" for value in editable):
                continue
            sku = str(values[1]).strip().upper() if values[1] is not None else ""
            name = str(values[2]).strip() if values[2] is not None else ""
            if not sku:
                raise ValueError(f"La fila {row_num} necesita el SKU del producto")
            if len(sku) > 100:
                raise ValueError(f"El SKU de la fila {row_num} supera los 100 caracteres")
            if len(name) > 255:
                raise ValueError(f"El nombre de la fila {row_num} supera los 255 caracteres")
            if sku in seen_skus:
                raise ValueError(f"El SKU {sku} está repetido en la plantilla (fila {row_num})")
            seen_skus.add(sku)

            price = parse_amount(values[5], "precio de venta", row_num, required=False, allow_zero=False)
            quantity_decimal = parse_amount(values[6], "cantidad recibida", row_num, required=True, allow_zero=False)
            if quantity_decimal != quantity_decimal.to_integral_value() or quantity_decimal > 2147483647:
                raise ValueError(f"La cantidad recibida de la fila {row_num} debe ser un entero positivo dentro del rango permitido")
            if price is not None and price > Decimal("9999999999.99"):
                raise ValueError(f"El precio de la fila {row_num} excede el máximo admitido")
            cost = parse_amount(values[7], "costo unitario", row_num)
            if cost is not None and cost > Decimal("99999999.999999"):
                raise ValueError(f"El costo de la fila {row_num} excede el máximo admitido")
            critical_decimal = parse_amount(values[8], "stock crítico", row_num)
            if critical_decimal is not None and (critical_decimal != critical_decimal.to_integral_value() or critical_decimal > 2147483647):
                raise ValueError(f"El stock crítico de la fila {row_num} debe ser un entero dentro del rango permitido")
            image_url = str(values[9]).strip() if values[9] is not None and str(values[9]).strip() else None
            if image_url and len(image_url) > 2000:
                raise ValueError(f"La URL de imagen de la fila {row_num} es demasiado larga")
            rows.append({
                "row_num": row_num,
                "sku": sku,
                "nombre": name,
                "descripcion": str(values[3]).strip() if values[3] is not None and str(values[3]).strip() else None,
                "tamano": str(values[4]).strip() if values[4] is not None and str(values[4]).strip() else None,
                "precio": price,
                "stock_inicial": int(quantity_decimal),
                "costo_nio": cost,
                "stock_critico": int(critical_decimal) if critical_decimal is not None else 5,
                "imagen_url": image_url,
            })
        if not rows:
            raise ValueError("No hay productos para importar")

        # Bloquear productos existentes evita que dos recepciones simultáneas
        # pierdan unidades al actualizar el kardex.
        existing_products = {
            (product.sku or "").upper(): product
            for product in db.query(Producto)
            .filter(func.upper(Producto.sku).in_([item["sku"] for item in rows]))
            .with_for_update()
            .all()
        }
        for item in rows:
            product = existing_products.get(item["sku"])
            if product and product.marca_id != marca_id:
                raise ValueError(
                    f"El SKU {item['sku']} pertenece a otra marca (fila {item['row_num']})."
                )
            if not product:
                if not item["nombre"]:
                    raise ValueError(f"La fila {item['row_num']} necesita el nombre para crear un producto nuevo")
                if item["precio"] is None:
                    raise ValueError(f"La fila {item['row_num']} necesita el precio para crear un producto nuevo")

        try:
            ingreso = Ingreso(
                marca_id=marca_id,
                usuario_id=usuario_id,
                tipo_carga="IMPORTACION EXCEL",
                observaciones=filename or "Importación de plantilla de productos",
                estado_id=1,
                fecha_ingreso=datetime.utcnow(),
            )
            db.add(ingreso)
            db.flush()

            productos_procesados = []
            productos_creados = 0
            productos_actualizados = 0
            for item in rows:
                producto = existing_products.get(item["sku"])
                if producto:
                    productos_actualizados += 1
                else:
                    producto = Producto(
                        marca_id=marca_id,
                        sku=item["sku"],
                        nombre=item["nombre"],
                        descripcion=item["descripcion"],
                        tamano=item["tamano"],
                        precio=item["precio"],
                        costo_nio=item["costo_nio"],
                        stock_inicial=0,
                        stock_critico=item["stock_critico"],
                        imagen_url=item["imagen_url"],
                        estado_id=1,
                        codigo_qr=f"BHPOS-{item['sku']}",
                    )
                    db.add(producto)
                    productos_creados += 1
                productos_procesados.append((producto, item))

            db.flush()
            detalles, movimientos = [], []
            for producto, item in productos_procesados:
                cantidad = item["stock_inicial"]
                detalles.append(IngresoDetalle(
                    ingreso_id=ingreso.id,
                    producto_id=producto.id,
                    cantidad=cantidad,
                    precio_ingreso=item["costo_nio"],
                ))
                movimientos.append(Kardex(
                    producto_id=producto.id,
                    usuario_id=usuario_id,
                    tipo_movimiento="INGRESO",
                    cantidad=cantidad,
                    referencia_id=ingreso.id,
                    motivo=f"Importación de productos desde {ingreso.observaciones}",
                ))
            db.add_all(detalles + movimientos)
            db.commit()
        except Exception:
            db.rollback()
            raise
        return {
            "ingreso_id": ingreso.id,
            "creados": productos_creados,
            "actualizados": productos_actualizados,
            "unidades": sum(item["stock_inicial"] for item in rows),
        }

    @staticmethod
    def delete(db: Session, marca_id: int) -> bool:
        marca = db.query(Marca).filter(Marca.id == marca_id).first()
        if not marca:
            return False

        # Eliminar cuentas asociadas
        db.query(Cuenta).filter(Cuenta.marca_id == marca_id).delete()
        db.delete(marca)
        db.commit()
        return True
